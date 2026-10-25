from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from copy import deepcopy
from pathlib import Path
P=Path(__file__).parent;src=Path(r'C:\Users\aaron\Downloads\25BCE2640 ESA LAB8 ASSESSMENTv2.docx')
d=Document(src);b=d._element.body;els=list(b)
def txt(e):return ''.join(e.itertext()) if False else ''.join(x.text or '' for x in e.iter(qn('w:t')))
def settext(e,s):
 ts=list(e.iter(qn('w:t')))
 if ts:
  ts[0].text=s
  for t in ts[1:]:t.text=''
def clonep(index,s='',after=None):
 e=deepcopy(els[index]);pp=e.find(qn('w:pPr'))
 for c in list(e):
  if c.tag!=qn('w:pPr'):e.remove(c)
 r=OxmlElement('w:r');oldr=els[index].find(qn('w:r'));rp=oldr.find(qn('w:rPr')) if oldr is not None else None
 if rp is not None:r.append(deepcopy(rp))
 t=OxmlElement('w:t');t.text=s;r.append(t);e.append(r)
 if after is not None:
  if pp is None:pp=OxmlElement('w:pPr');e.insert(0,pp)
  sp=pp.find(qn('w:spacing'))
  if sp is None:sp=OxmlElement('w:spacing');pp.append(sp)
  sp.set(qn('w:after'),str(after*20))
 return e
def pagebreak(e):
 pp=e.find(qn('w:pPr'))
 if pp is None:pp=OxmlElement('w:pPr');e.insert(0,pp)
 x=OxmlElement('w:pageBreakBefore');x.set(qn('w:val'),'1');pp.append(x)
def add(e):b.insert(len(b)-1,e);return e
def p(i,s='',after=None):return add(clonep(i,s,after))
title='Four-Way Traffic Light Control and LED Fading Using Arduino Uno'
# Preserve original cover, contents and signature box; replace all experiment pages.
for e in list(b)[9:-1]:b.remove(e)
settext(els[1],'Assessment 9: '+title)
for cell in d.tables[0]._tbl.iter(qn('w:tc')):
 s=txt(cell)
 if s.strip()=='8':settext(cell,'9')
 if 'Assessment no:' in s:settext(cell,'Assessment no: 9')
 if 'Interfacing LED, Servo Motor' in s:settext(cell,title)
 if 'Assigned date of Task' in s:settext(cell,'Assigned date of Task:')
 if 'Date of Completion' in s:settext(cell,'Date of Completion:')
 if 'VTOP uploading Date and time' in s:settext(cell,'VTOP uploading Date and time:')
# Reuse the source connection table's borders, cell widths and typography.
def table(rows):
 t=deepcopy(els[25]);trs=t.findall(qn('w:tr'));header=deepcopy(trs[0]);sample=deepcopy(trs[1])
 for row in trs:t.remove(row)
 for j,vals in enumerate(rows):
  row=deepcopy(header if j==0 else sample)
  for cell,value in zip(row.findall(qn('w:tc')),vals):settext(cell,value)
  t.append(row)
 return add(t)
def blank_section(heading,index,height):
 h=p(index,heading);pagebreak(h)
 # A genuinely empty paragraph reserves editable space without placeholder text.
 q=p(12,'',height);pp=q.find(qn('w:pPr'))
 for k in ['keepNext','keepLines']:
  x=pp.find(qn('w:'+k))
  if x is not None:pp.remove(x)
def result(text,inference):
 h=p(55,'Result');pagebreak(h);p(56,text);p(57,'Inference');p(58,inference)
def start(num,question,aim,tools,theory):
 h=p(9,'Program '+str(num));pagebreak(h);p(10,question);p(11,'Aim of the Experiment');p(12,aim);p(13,'Tools Required')
 for tool in tools:p(14,tool)
 p(21,'Brief Theory')
 for paragraph in theory:p(22,paragraph)
start(1,'Design a hardware circuit for a four-way traffic-light junction. Provide red, yellow and green LEDs for each direction and use 220 ohm resistors.',
'To interface twelve LEDs with Arduino Uno and operate a timed four-way traffic-light sequence that gives one approach a green signal while the others remain red.',
['Arduino Uno R3','Four red LEDs, four yellow LEDs and four green LEDs','Twelve 220 ohm resistors, one for each LED','Breadboard','Jumper wires','USB cable','Computer with Arduino IDE'],
['A traffic controller assigns right of way through red, yellow and green signals. Red means stop, green permits movement and yellow signals the end of the current phase. This model serves North, East, South and West one at a time.',
'Twelve digital outputs independently drive twelve LEDs through individual 220 ohm series resistors. The proposed cycle uses 5 seconds of green, 2 seconds of yellow and 1 second with all directions red before the next approach is released. These timings are selected for the demonstration, not specified by the question.'])
h=p(24,'Circuit Connections');pagebreak(h)
rows=[('Component terminal','Connection')]
for direction,pins in [('North',(2,3,4)),('East',(5,6,7)),('South',(8,9,10)),('West',(11,12,13))]:
 for color,pin in zip(['red','yellow','green'],pins):rows.append((direction+' '+color+' LED anode','Arduino D'+str(pin)+' through its own 220 ohm resistor'))
rows.extend([('All LED cathodes','Common ground rail connected to Arduino GND'),('Arduino USB port','Computer USB port')]);table(rows)
p(12,'Wire each approach in red-yellow-green order. Use only one LED per output pin, check every LED polarity and switch off power before changing connections.')
blank_section('Arduino Code with Comments',27,280)
p(41,'Manual calculation:');p(42,'None');p(43,'Procedure')
for i,s in enumerate(['Label the four approaches North, East, South and West and arrange three LEDs at each approach.','Disconnect power and wire each LED anode through a separate 220 ohm resistor to the pin listed in the connection table. Connect every cathode to GND.','Connect the Uno to the computer, select Arduino Uno and the correct port, and prepare the traffic-control sketch in the blank code section.','Configure D2 to D13 as outputs. Begin with all red LEDs ON and all yellow and green LEDs OFF.','For North, turn its red OFF and green ON for 5 seconds; keep East, South and West red.','Turn North green OFF and yellow ON for 2 seconds, then yellow OFF and red ON. Hold all four red signals for 1 second.','Repeat the same sequence for East, South and West, then repeat the complete cycle.','Verify and upload the completed sketch. Check that two approaches never show green together and record the actual observations.'],1):p(44,f'{i}. {s}')
blank_section('Output Snips',50,400)
result('Expected result: North, East, South and West should receive green signals in turn. Each green phase should be followed by yellow and an all-red interval. All other approaches should remain red while one approach is active. Confirm the sequence on the physical circuit.',
'Arduino digital outputs can coordinate a multi-direction traffic-light sequence. Individual current-limiting resistors protect the LEDs, while a shared ground completes each circuit. The all-red interval separates successive phases and helps demonstrate how signal interlocking avoids conflicting green indications.')
start(2,'Design an Arduino circuit that fades an LED using a PWM pin. Vary brightness from 0 to 255 and back to 0 using brightness and fadeAmount, and observe the effect of changing the delay.',
'To control LED brightness through Arduino PWM and produce a repeating fade-in and fade-out cycle by changing the duty cycle.',
['Arduino Uno R3','One LED','One 220 ohm resistor','Breadboard','Jumper wires','USB cable','Computer with Arduino IDE'],
['Pulse Width Modulation switches a digital output rapidly between LOW and HIGH. Changing the fraction of time spent HIGH changes the average light output. On the Uno, analogWrite() accepts a value from 0 to 255: 0 keeps the LED off and 255 keeps it fully on.',
'Use PWM pin D9 and initialise brightness to 0 and fadeAmount to 5. After writing the brightness, change it by fadeAmount. Reverse the sign of fadeAmount at 0 or 255. A short delay between updates controls the apparent fading speed; a longer delay makes the fade slower.'])
h=p(24,'Circuit Connections');pagebreak(h);table([('Component terminal','Connection'),('Arduino D9 (PWM)','220 ohm resistor, then LED anode (long lead)'),('LED cathode (short lead)','Arduino GND through the common ground rail'),('Arduino USB port','Computer USB port')])
p(12,'Use D9 because it supports PWM. A series resistor is required; do not connect the LED directly between the output pin and GND.')
blank_section('Arduino Code with Comments',27,280)
p(41,'Manual calculation:');p(42,'None');p(43,'Procedure')
for i,s in enumerate(['Disconnect power. Connect D9 through the 220 ohm resistor to the LED anode and connect the cathode to Arduino GND.','Connect the board by USB, select Arduino Uno and the correct port, and prepare the fading sketch in the blank code section.','Declare brightness with an initial value of 0 and fadeAmount with a value of 5. Use D9 as the LED output.','Write brightness to D9 with analogWrite(), then add fadeAmount to brightness on each loop.','When brightness reaches 0 or 255, reverse the sign of fadeAmount so the next updates move in the opposite direction. Keep the value within 0 to 255.','Begin with a delay of 30 milliseconds between brightness updates. Verify and upload the completed sketch.','Observe the LED gradually becoming brighter and dimmer through several cycles.','Repeat with shorter and longer delays, such as 10 and 60 milliseconds, and compare the fading speeds. Record actual observations.'],1):p(44,f'{i}. {s}')
blank_section('Output Snips',50,400)
result('Expected result: The LED should gradually brighten from OFF to full brightness, then dim back to OFF repeatedly. Reversing fadeAmount at the limits should reverse the direction of fading. Increasing the update delay should slow the cycle; reducing it should speed up the effect.',
'PWM controls perceived LED brightness by varying duty cycle rather than generating a continuously variable output voltage. The brightness variable determines the PWM level, while fadeAmount determines the step direction and size. The delay between updates controls how quickly the brightness changes.')
# Retain the final assessment table and the user's details.
report=deepcopy(els[261]);pagebreak(report);add(report);last=deepcopy(els[262]);add(last)
for cell in last.iter(qn('w:tc')):
 s=txt(cell)
 if 'Title of Assessment:' in s:settext(cell,'Title of Assessment: '+title)
 if 'Assessment no.' in s:settext(cell,'Assessment no.: 9')
 if 'September' in s or 'October' in s:settext(cell,'')
add(clonep(263,''))
# Update all header/footer assessment labels and the first-page title.
for section in d.sections:
 for hf in [section.header,section.first_page_header,section.even_page_header,section.footer,section.first_page_footer,section.even_page_footer]:
  for para in hf.paragraphs:
   s=para.text
   if 'Assessment' in s:
    s=s.replace('Assessment: 8','Assessment: 9').replace('Assessment 8','Assessment 9')
    old='Interfacing LED, Servo Motor, Ultrasonic Distance Sensor and LCD Hardware with Arduino Uno'
    s=s.replace(old,title)
    if s!=para.text:settext(para._p,s)
# Contents-page references verified against the rendered fifteen-page draft.
contents=d.tables[1]
page_refs=['3, 9','3, 9','3, 9','3, 9','4, 10','5, 11','6, 12','6, 12','7, 13','8, 14','8, 14']
for ri,row in enumerate(contents.rows):
 if ri==1:settext(row.cells[0]._tc,'Program 1, 2')
 if ri>0:settext(row.cells[1]._tc,page_refs[ri-1])
out=P/'25BCE2640_ESA_Assessment_9.docx';d.save(out)
print(out)
