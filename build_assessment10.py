from pathlib import Path
P=Path(__file__).parent
base=(P/'build_assessment9.py').read_text()
prefix=base.split("start(1,")[0]
prefix=prefix.replace("title='Four-Way Traffic Light Control and LED Fading Using Arduino Uno'", "title='Seven Segment Display Interfacing and Distance Measurement Using Arduino Uno'")
prefix=prefix.replace('Assessment 9:', 'Assessment 10:').replace("settext(cell,'9')", "settext(cell,'10')").replace('Assessment no: 9','Assessment no: 10')
exec(compile(prefix,'template_helpers','exec'))

start(1,'Design a hardware circuit to interface a seven segment display with Arduino Uno. Display 0 to 9 and A to F, using 220 ohm resistors.',
'To interface a common-cathode seven segment display with Arduino Uno and display the sixteen hexadecimal symbols from 0 to F in sequence.',
['Arduino Uno R3','One common-cathode seven segment display','Seven 220 ohm resistors, one per segment','Breadboard and jumper wires','USB cable','Computer with Arduino IDE'],
['A seven segment display contains LED segments labelled a to g. Different combinations represent decimal digits and hexadecimal symbols. For a common-cathode display, connect the common terminals to GND; a HIGH output lights a segment through its own 220 ohm resistor.',
'Assign segments a to g to D2 to D8 in order. A sixteen-entry lookup table defines the segment patterns for 0 to 9 and A to F. Show each pattern for about one second before advancing. Lowercase b and d are used because uppercase B and D cannot be represented distinctly on seven segments.'])
p(24,'Circuit Connections')
table([('Component terminal','Connection')]+[(f'Display segment {s}',f'Arduino D{pin} through one 220 ohm resistor') for s,pin in zip('abcdefg',range(2,9))]+[('Common-cathode terminals','Arduino GND; verify the display datasheet pinout'),('Decimal point','Leave disconnected'),('Arduino USB port','Computer USB port')])
p(12,'Use segment labels rather than assuming physical pin numbers. Display pinouts vary by part. This connection plan is for a common-cathode display; a common-anode part requires different wiring and inverted logic.')
p(21,'Display Patterns')
table([('Symbol','Segments ON')]+list(zip(list('0123456789AbCdEF'),['a b c d e f','b c','a b d e g','a b c d g','b c f g','a c d f g','a c d e f g','a b c','a b c d e f g','a b c d f g','a b c e f g','c d e f g','a d e f','b c d e g','a d e f g','a e f g'])))
blank_section('Arduino Code with Comments',27,440)
p(41,'Manual calculation:');p(42,'None');p(43,'Procedure')
for i,s in enumerate(['Identify the common-cathode display and locate segments a to g using its datasheet. Disconnect power before wiring.','Connect the common cathodes to GND. Connect each segment to D2 to D8 through its own 220 ohm resistor, following the connection table.','Prepare the Arduino sketch in the reserved code section. Set the seven segment pins as outputs and initially turn every segment OFF.','Create patterns for all sixteen symbols using the Display Patterns table. Apply HIGH to active segments and LOW to inactive segments.','Display each symbol for about one second, then advance through 0 to 9 and A, b, C, d, E, F. Repeat the sequence continuously.','Select Arduino Uno and the correct port, verify the sketch and upload it.','Check each symbol against the pattern table. Correct wiring errors with power disconnected and capture clear output photographs.'],1):p(44,f'{i}. {s}')
blank_section('Output Snips',50,400)
result('Expected result: The display should show 0 to 9 followed by A, b, C, d, E and F, then repeat. Each symbol should use the correct segment pattern. Verify the sequence on the connected circuit.',
'Seven independently controlled LED segments can represent sixteen hexadecimal symbols using a lookup table. One resistor per segment limits current. Correct common-terminal wiring and segment mapping are essential for accurate characters.')

start(2,'Design a distance-measuring hardware circuit with six different LEDs and a 16 by 2 LCD. Show the measured distance on the LCD and switch ON the LED corresponding to that distance.',
'To measure distance using an HC-SR04 ultrasonic sensor, show the measured distance in centimetres on a 16 by 2 LCD and select one of six LEDs according to the measured range.',
['Arduino Uno R3','HC-SR04 ultrasonic distance sensor','16 by 2 LCD compatible with the LiquidCrystal library','10 kilohm potentiometer for LCD contrast','Six LEDs of different colours','Six 220 ohm resistors for the LEDs','220 ohm LCD backlight resistor if required by the module','Breadboard and jumper wires','USB cable and computer with Arduino IDE','Ruler or measuring tape and a flat target'],
['The HC-SR04 emits an ultrasonic burst after a 10 microsecond trigger pulse. Its ECHO pulse duration represents the outward and return travel time of the sound. Distance is calculated from this duration and displayed on the LCD; a ruler provides an independent comparison.',
'Use a parallel LCD in four-bit mode with the LiquidCrystal library. Compare each valid reading with six non-overlapping distance bands and light only the selected LED. The thresholds below are demonstration choices because the question does not specify them. If no valid echo is received, display No echo and turn all six LEDs OFF.'])
p(24,'Circuit Connections')
table([('Component terminal','Connection'),('HC-SR04 VCC and GND','Arduino 5 V and GND'),('HC-SR04 TRIG and ECHO','Arduino D2 and D3 respectively'),('LED 1 to LED 6 anodes','D4, D5, D6, D7, D8, D9 respectively; each through its own 220 ohm resistor'),('All six LED cathodes','Arduino GND'),('LCD VSS, VDD and RW','GND, 5 V and GND respectively'),('LCD VO contrast','10 kilohm potentiometer wiper; outer terminals to 5 V and GND'),('LCD RS and Enable','Arduino D10 and D11 respectively'),('LCD D4 and D5','Arduino D12 and D13 respectively'),('LCD D6 and D7','Arduino A0 and A1 used as digital outputs'),('LCD D0 to D3','Leave disconnected in four-bit mode'),('LCD backlight A and K','5 V through required current-limiting resistor, and GND; check module specifications'),('Arduino USB port','Computer USB port')])
p(12,'Use LiquidCrystal lcd(10, 11, 12, 13, A0, A1) for this pin plan. Join every component ground to Arduino GND. Label the six LEDs 1 to 6 so that their distance bands remain clear regardless of the colours available.')
p(21,'LED Distance Bands')
table([('Selected LED','Measured distance in cm'),('LED 1','2 to less than 10'),('LED 2','10 to less than 20'),('LED 3','20 to less than 30'),('LED 4','30 to less than 40'),('LED 5','40 to less than 50'),('LED 6','50 to 400'),('All LEDs OFF','No echo or reading outside 2 to 400')])
blank_section('Arduino Code with Comments',27,440)
p(41,'Manual calculation:')
p(42,'Distance in cm = echo duration in microseconds x 0.0343 / 2. The factor 0.0343 cm per microsecond approximates the speed of sound at room temperature; division by two accounts for the return journey.')
p(42,'Example only: for an echo duration of 1166 microseconds, distance = 1166 x 0.0343 / 2 = 20.00 cm. This falls in the LED 3 band. Compare measured readings with a ruler; this example is not an experimental observation.')
p(43,'Procedure')
for i,s in enumerate(['Disconnect power and wire the ultrasonic sensor, six LEDs and LCD according to the connection table. Use a separate resistor for each LED.','Connect USB power, adjust the LCD contrast potentiometer and initialise the LCD as 16 columns by 2 rows.','Prepare the sketch in the reserved section. Configure TRIG as OUTPUT, ECHO as INPUT and the six LED pins as OUTPUT. Start with all LEDs OFF.','Send a 10 microsecond trigger pulse. Measure the ECHO pulse with a bounded timeout, such as 30000 microseconds, and calculate distance in centimetres.','For each valid reading, turn all LEDs OFF and then turn ON only the LED for its distance band. Treat missing or out-of-range readings as invalid.','Show the distance on LCD row 1, for example Distance: 20.0cm, and the selected LED on row 2, for example LED: 3. Clear unused character positions when a value becomes shorter.','Verify and upload the sketch. Place a flat target at several ruler-measured positions, including both sides of each distance threshold, and allow about 100 milliseconds between readings.','Compare the displayed distance with the ruler reading. Check that exactly one LED lights for a valid distance, then capture the LCD and corresponding LED in the output photographs.'],1):p(44,f'{i}. {s}')
blank_section('Output Snips',50,400)
result('Expected result: The LCD should update the measured distance in centimetres, and exactly one of six LEDs should indicate its band. Moving the target across a threshold should change the selected LED. An invalid reading should show No echo or Out of range with all LEDs OFF.',
'Ultrasonic echo timing provides a distance estimate, while the LCD and six LEDs present numeric and range-based outputs. Accuracy depends on target alignment, surface and ambient conditions. Comparing readings with a ruler helps assess the measurement error.')

report=deepcopy(els[261]);pagebreak(report);add(report);last=deepcopy(els[262]);add(last)
for cell in last.iter(qn('w:tc')):
 s=txt(cell)
 if 'Title of Assessment:' in s:settext(cell,'Title of Assessment: '+title)
 if 'Assessment no.' in s:settext(cell,'Assessment no.: 10')
 if 'September' in s or 'October' in s:settext(cell,'')
add(clonep(263,''))
for section in d.sections:
 for hf in [section.header,section.first_page_header,section.even_page_header,section.footer,section.first_page_footer,section.even_page_footer]:
  for t in hf._element.iter(qn('w:t')):
   if t.text:t.text=t.text.replace('Assessment: 8','Assessment: 10').replace('Assessment 8','Assessment 10').replace('Lab Serial no.: 8','Lab Serial no.: 10')
# Match the compact heading spacing while retaining editable blank sections.
from docx.text.paragraph import Paragraph
heads={'Aim of the Experiment','Tools Required','Brief Theory','Circuit Connections','Display Patterns','LED Distance Bands','Manual calculation:','Procedure','Result','Inference'}
for para in d.paragraphs:
 f=para.paragraph_format;t=para.text.strip()
 if t in heads:
  f.space_before=Pt(0);f.space_after=Pt(0);f.keep_with_next=True
  if t!='Result':f.page_break_before=False
 if t.startswith(tuple(f'{i}. ' for i in range(1,10))):f.keep_with_next=False
 if t=='None':f.space_after=Pt(0);f.keep_with_next=False
 # Convert the tools to compact bullets, matching the assessment template.
 if t in ['Tools Required']:continue
for heading in [p for p in d.paragraphs if p.text.strip()=='Tools Required']:
 el=heading._p.getnext()
 while el is not None and txt(el).strip()!='Brief Theory':
  para=Paragraph(el,heading._parent);para.paragraph_format.space_before=Pt(0);para.paragraph_format.space_after=Pt(2)
  if not para.text.startswith('• '):settext(el,'• '+para.text)
  para.paragraph_format.keep_with_next=False
  el=el.getnext()
contents=d.tables[1]
settext(contents.rows[1].cells[0]._tc,'Program 1, 2')
for row in contents.rows[1:]:settext(row.cells[1]._tc,'')
from docx.enum.text import WD_TAB_ALIGNMENT
from docx.shared import Inches
for section in d.sections:
 for hf in [section.header,section.even_page_header]:
  for pi,text in enumerate(['Embedded Systems Architecture Assessment: 10\tLab Serial no.: 10','Name: Aaron Robertson\tReg. No.: 25BCE2640']):
   while len(hf.paragraphs)<=pi:hf.add_paragraph()
   para=hf.paragraphs[pi];para.clear();para.paragraph_format.tab_stops.clear_all()
   para.paragraph_format.tab_stops.add_tab_stop(Inches(6.5),WD_TAB_ALIGNMENT.RIGHT)
   r=para.add_run(text);r.font.name='Calibri';r.font.size=Pt(10);r.bold=True
 for hf in [section.first_page_footer,section.footer,section.even_page_footer]:
  for para in hf.paragraphs:
   if 'Assessment 8' in para.text:settext(para._p,para.text.replace('Assessment 8','Assessment 10'))
for para in d.paragraphs:
 if para.text.strip() in ['Circuit Connections','Manual calculation:']:para.paragraph_format.page_break_before=True
 if para.text.startswith('• '):
  num=para._p.find('w:pPr/w:numPr',para._p.nsmap)
  if num is not None:num.getparent().remove(num)
for tab in d.tables[3:-1]:
 for row in tab.rows:
  row.height=None
  for cell in row.cells:
   for para in cell.paragraphs:
    para.paragraph_format.space_before=Pt(0);para.paragraph_format.space_after=Pt(0);para.paragraph_format.line_spacing=1
    for r in para.runs:r.font.size=Pt(11)
out=P/'25BCE2640_ESA_Assessment_10.docx';d.save(out);print(out)
