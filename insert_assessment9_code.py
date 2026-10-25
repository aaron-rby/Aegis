from pathlib import Path
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_BREAK
P=Path(__file__).parent
src=P/'25BCE2640_ESA_Assessment_9_compact.docx';d=Document(src)
files=[P/'assessment9_code/Traffic_Light/Traffic_Light.ino',P/'assessment9_code/LED_Fading/LED_Fading.ino']
headings=[p for p in d.paragraphs if p.text.strip()=='Arduino Code with Comments']
assert len(headings)==2
for heading,file in zip(headings,files):
    blank=heading._p.getnext()
    from docx.text.paragraph import Paragraph
    p=Paragraph(blank,heading._parent)
    assert not p.text.strip()
    p.clear();p.style=d.styles['Normal']
    fmt=p.paragraph_format;fmt.space_before=Pt(0);fmt.space_after=Pt(6);fmt.line_spacing=Pt(10.5)
    fmt.keep_with_next=False;fmt.keep_together=False;fmt.page_break_before=False
    fmt.left_indent=Pt(0);fmt.right_indent=Pt(0);fmt.first_line_indent=Pt(0)
    run=p.add_run(file.read_text().rstrip())
    run.font.name='Courier New';run.font.size=Pt(9.5);run.bold=False
out=P/'25BCE2640_ESA_Assessment_9_with_Code.docx';d.save(out)
# Confirm both complete sketches are present and output sections stay empty.
check=Document(out)
for file in files:assert file.read_text().rstrip() in '\n'.join(p.text for p in check.paragraphs)
for i,p in enumerate(check.paragraphs):
    if p.text.strip()=='Output Snips':assert not check.paragraphs[i+1].text.strip()
print(out)
print('Inserted both sketches; existing connection tables and blank Output Snips preserved.')
