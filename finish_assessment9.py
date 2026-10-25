from pathlib import Path
import re, urllib.request, urllib.parse
from lxml import html
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.text.paragraph import Paragraph
from docx.oxml import OxmlElement

P=Path(__file__).parent
d=Document(P/'25BCE2640_ESA_Assessment_9_with_Code.docx')
files=[P/'assessment9_code/Traffic_Light/Traffic_Light.ino',P/'assessment9_code/LED_Fading/LED_Fading.ino']
for idx,(heading,file) in enumerate(zip([p for p in d.paragraphs if p.text.strip()=='Arduino Code with Comments'],files)):
    lines=file.read_text().splitlines()
    code='\n'.join(re.sub(r'\s*//.*$','',line) for line in lines if not line.lstrip().startswith('//')).strip()
    code=re.sub(r'\n{3,}','\n\n',code)
    code=('// North, East, South, West pin order.\n' if idx==0 else '// Fade the LED using PWM on pin 9.\n')+code
    # Use the HTML produced by hilite.me, retaining its token colours in Word.
    data=urllib.parse.urlencode({'code':code,'lexer':'arduino','style':'colorful'}).encode()
    result=urllib.request.urlopen('https://hilite.me/api',data,timeout=45).read().decode()
    (P/f'assessment9_hilite_{idx+1}.html').write_text(result,encoding='utf-8')
    pre=html.fromstring(result).xpath('//pre')[0]
    p=Paragraph(heading._p.getnext(),heading._parent);p.clear()
    fmt=p.paragraph_format;fmt.line_spacing=Pt(10.5);fmt.space_before=Pt(0);fmt.space_after=Pt(6)
    fmt.keep_together=False;fmt.keep_with_next=False
    def add(text,style=''):
        if not text:return
        r=p.add_run(text);r.font.name='Courier New';r.font.size=Pt(9.5);r.bold='bold' in style;r.italic='italic' in style
        color=re.search(r'(?<!-)color:\s*#([0-9a-fA-F]{6})',style)
        r.font.color.rgb=RGBColor.from_string(color.group(1) if color else '000000')
    add(pre.text)
    for el in pre:
        add(el.text,el.get('style',''));add(el.tail)

shots=[
 [Path('C:/Users/aaron/Downloads/WhatsApp Image 2026-10-22 at 10.58.55 AM (1).jpeg'),Path('C:/Users/aaron/Downloads/WhatsApp Image 2026-10-22 at 10.58.56 AM (1).jpeg')],
 [Path('C:/Users/aaron/AppData/Local/Temp/codex-clipboard-a119626f-5f3d-40cb-bd19-6bdc561d64e8.png')]
]
captions=[['Traffic-light circuit with the red LED illuminated.','Traffic-light circuit showing a different LED state.'],['Single-LED circuit used for the PWM fading experiment.']]
for idx,heading in enumerate([p for p in d.paragraphs if p.text.strip()=='Output Snips']):
    blank=heading._p.getnext();blank.getparent().remove(blank)
    # Keep results on their original separate page.
    result=heading._p.getnext();Paragraph(result,heading._parent).paragraph_format.page_break_before=True
    anchor=heading._p
    for photo,caption in zip(shots[idx],captions[idx]):
        el=OxmlElement('w:p');anchor.addnext(el);p=Paragraph(el,heading._parent)
        p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_after=Pt(3)
        p.paragraph_format.space_before=Pt(0);p.paragraph_format.line_spacing=1
        p.paragraph_format.keep_with_next=True
        p.add_run().add_picture(str(photo),height=Inches(3.45 if idx==0 else 5.9))
        cap=OxmlElement('w:p');el.addnext(cap);cp=Paragraph(cap,heading._parent)
        cp.alignment=WD_ALIGN_PARAGRAPH.CENTER;cp.paragraph_format.space_after=Pt(8)
        cp.paragraph_format.space_before=Pt(0);cp.paragraph_format.line_spacing=1
        r=cp.add_run(caption);r.font.name='Times New Roman';r.font.size=Pt(11)
        anchor=cap
out=P/'25BCE2640_ESA_Assessment_9_Final.docx';d.save(out)
print(out)
